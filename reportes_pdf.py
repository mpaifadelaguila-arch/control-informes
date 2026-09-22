"""
reportes_pdf.py — genera, sin ninguna llamada a IA, los PDF que alimentan los
Anexos del informe (anexos.py):

    generar_pdf_checklist_vt_por_tag(...) -> Anexo B: exporta cada hoja
        (línea) del VT-CHECK LIST ya parchado (checklist.py) tal cual, con
        su formato, colores y fotografías originales, vía LibreOffice --
        el Anexo B muestra el checklist real que llenó el inspector, no un
        resumen reconstruido.

    generar_pdf_psaim_por_tag(...)      -> Anexo C: exporta la hoja real del
        PSAIM (Siemens Energy) de cada línea tal cual -- con su formato,
        logo y tabla de TML originales, vía LibreOffice -- el Anexo C
        muestra el reporte de ultrasonido real, no un resumen reconstruido
        (mismo criterio que el Anexo B del checklist VT).
"""
import os
import shutil
import tempfile

import openpyxl

import anexos
from checklist import CELDA_LINEA, extraer_hoja_a_xlsx


def generar_pdf_checklist_vt_por_tag(ruta_checklist_parchado, dir_salida):
    """Por cada hoja (línea) del VT-CHECK LIST ya parchado, exporta esa
    hoja sola -- con su formato, colores y fotografías originales
    intactos -- a un PDF independiente vía LibreOffice (checklist.
    extraer_hoja_a_xlsx aísla la hoja a nivel de XML crudo, igual que el
    parchado, para no perder ninguna fotografía en el camino)."""
    os.makedirs(dir_salida, exist_ok=True)
    wb = openpyxl.load_workbook(ruta_checklist_parchado)

    rutas = {}
    with tempfile.TemporaryDirectory() as tmp_hojas:
        for i, nombre_hoja in enumerate(wb.sheetnames):
            ws = wb[nombre_hoja]
            tag_val = ws[CELDA_LINEA].value
            if not tag_val:
                continue
            tag = str(tag_val).strip()

            ruta_hoja_xlsx = os.path.join(tmp_hojas, f"hoja_{i}.xlsx")
            extraer_hoja_a_xlsx(ruta_checklist_parchado, nombre_hoja, ruta_hoja_xlsx)

            ruta_pdf_generado = anexos.convertir_xlsx_a_pdf(ruta_hoja_xlsx, tmp_hojas)
            if not ruta_pdf_generado:
                continue

            ruta_pdf_final = os.path.join(dir_salida, f"Checklist_VT_{_slug(tag)}.pdf")
            shutil.copyfile(ruta_pdf_generado, ruta_pdf_final)
            rutas[tag] = ruta_pdf_final

    return rutas


def _slug(texto):
    return "".join(c if c.isalnum() else "_" for c in str(texto)).strip("_")


def generar_pdf_psaim_por_tag(psaim_por_tag, dir_salida):
    """Por cada línea con PSAIM, exporta la hoja REAL de ese archivo (la que
    quedó identificada al emparejar el PSAIM con la línea, ver
    generar_informe._leer_psaim_por_linea) tal cual -- con su formato, logo
    y tabla de TML originales intactos -- a un PDF independiente vía
    LibreOffice, igual que el Anexo B del checklist VT (checklist.
    extraer_hoja_a_xlsx aísla la hoja a nivel de XML crudo, sin perder
    formato en el camino)."""
    os.makedirs(dir_salida, exist_ok=True)
    rutas = {}

    with tempfile.TemporaryDirectory() as tmp_hojas:
        for tag, datos in psaim_por_tag.items():
            ruta_archivo = datos.get("ruta_archivo")
            nombre_hoja = datos.get("nombre_hoja")
            if not ruta_archivo or not nombre_hoja:
                continue

            ruta_hoja_xlsx = os.path.join(tmp_hojas, f"psaim_{_slug(tag)}.xlsx")
            extraer_hoja_a_xlsx(ruta_archivo, nombre_hoja, ruta_hoja_xlsx)

            ruta_pdf_generado = anexos.convertir_xlsx_a_pdf(ruta_hoja_xlsx, tmp_hojas)
            if not ruta_pdf_generado:
                continue

            ruta_pdf_final = os.path.join(dir_salida, f"PSAIM_{_slug(tag)}.pdf")
            shutil.copyfile(ruta_pdf_generado, ruta_pdf_final)
            rutas[tag] = ruta_pdf_final

    return rutas
