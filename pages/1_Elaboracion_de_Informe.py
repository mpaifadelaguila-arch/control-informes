import io
import os
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import streamlit as st
from docx import Document

# Configuración inicial de rutas
RUTA_ACTUAL = Path(__file__).resolve().parent
DIR_RAIZ = RUTA_ACTUAL.parent if RUTA_ACTUAL.name == "pages" else RUTA_ACTUAL

if str(DIR_RAIZ) not in sys.path:
    sys.path.insert(0, str(DIR_RAIZ))

RUTA_PLANTILLA = DIR_RAIZ / "plantilla_base.docx"
RUTA_MAESTRA = DIR_RAIZ / "BASE_DE_DATOS_DE_LINEAS_FASE1.xlsx"

st.set_page_config(
    page_title="Elaboración de Informes - Ademinsac",
    page_icon="📋",
    layout="wide"
)

st.title("Módulo de Elaboración de Informes")
st.markdown("---")

# Verificación visual rápida de recursos
c1, c2 = st.columns(2)
with c1:
    if RUTA_MAESTRA.exists():
        st.success("🟢 Base de Datos Maestra detectada en la raíz")
    else:
        st.warning("⚠️ No se encontró la Base Maestra en la raíz")
with c2:
    if RUTA_PLANTILLA.exists():
        st.success("🟢 Plantilla Word base detectada")
    else:
        st.warning("⚠️ Falta plantilla_base.docx en la raíz")

st.markdown("---")
st.subheader("📁 Carga de Archivos para el Informe")

col1, col2, col3 = st.columns(3)
with col1:
    f_m3m6 = st.file_uploader("1. Archivo M3 y M6 (Excel)", type=["xlsx", "xls"], key="m3m6")
with col2:
    f_cons = st.file_uploader("2. Archivo Consolidadas (Excel)", type=["xlsx", "xls"], key="cons")
with col3:
    f_fotos = st.file_uploader("3. Fotos de la Unidad", type=["jpg", "jpeg", "png"], accept_multiple_files=True, key="fotos")

st.markdown("---")

if st.button("🚀 Procesar y Generar Documentos", type="primary", use_container_width=True):
    if not (f_m3m6 and f_cons and f_fotos):
        st.error("Por favor, asegúrate de cargar los archivos obligatorios (M3/M6, Consolidadas y Fotos).")
    else:
        with st.spinner("Generando entregables..."):
            try:
                # Lectura de archivos cargados
                df_m3 = pd.read_excel(f_m3m6)
                df_c = pd.read_excel(f_cons)

                # Generación de Word básico seguro
                doc = Document(RUTA_PLANTILLA) if RUTA_PLANTILLA.exists() else Document()
                if not RUTA_PLANTILLA.exists():
                    doc.add_heading("INFORME TÉCNICO DE INTEGRIDAD", 0)

                out_word = io.BytesIO()
                doc.save(out_word)
                
                # Generación de Excel
                out_excel = io.BytesIO()
                wb = openpyxl.load_workbook(f_cons)
                wb.save(out_excel)

                # Guardar en session_state para los botones de descarga
                st.session_state["res_word"] = out_word.getvalue()
                st.session_state["res_excel"] = out_excel.getvalue()
                st.session_state["ok_generado"] = True

                st.success("¡Informes generados exitosamente!")
            except Exception as ex:
                st.error(f"Ocurrió un error al procesar los archivos: {str(ex)}")

if st.session_state.get("ok_generado", False):
    st.markdown("---")
    st.subheader("📥 Descarga de Archivos")
    b_col1, b_col2 = st.columns(2)
    with b_col1:
        st.download_button(
            "📄 Descargar Informe Word",
            data=st.session_state["res_word"],
            file_name=f"Informe_{datetime.now():%Y%m%d_%H%M%S}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
    with b_col2:
        st.download_button(
            "📊 Descargar Excel Consolidado",
            data=st.session_state["res_excel"],
            file_name=f"Checklist_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
