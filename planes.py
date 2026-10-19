"""Lógica de la matriz de planes especiales (sin dependencias de Streamlit).

La matriz de planes tiene las mismas 14 columnas que la matriz del plan general
de tuberías, pero es independiente. El plan se identifica en la columna
"ALCANCE DEL SERVICIO" (puntos de inyección, dead legs o sulfidación).

Un informe es la unidad de control: se identifica por mes + código de informe,
o por mes + grupo cuando el código aún está por asignar (igual que en la matriz
general).
"""

import io
import json
import os
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COL_ESTADO = "ESTADO - ELABORACIÓN "
COLUMNAS = [
    "ITEM POR MES",
    "IT2",
    "UNIDAD",
    "MES",
    "LINEAS",
    "CODIGO DE INFORME",
    "GRUPO DE TUBERÍAS",
    "SAP",
    "ALCANCE DEL SERVICIO",
    "NOTAS",
    COL_ESTADO,
    "RESPONSABLE",
    "OBSERVACIÓN",
    "VALORIZACIÓN",
]
ORDEN_MESES = [
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
    "AGOSTO", "SETIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
]
PLANES = ["PUNTOS DE INYECCIÓN", "DEAD LEG", "SULFIDACIÓN"]

# Estados de un informe (mutuamente excluyentes).
FINALIZADO = "Finalizado"
PEND_INSPECCION = "Pendiente inspección"
PEND_ELABORACION = "Pendiente elaboración"
OTRO = "En proceso"


def texto_normalizado(valor):
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return ""
    texto = str(valor).strip().upper()
    return texto.translate(str.maketrans("ÁÉÍÓÚÜÑ", "AEIOUUN"))


def texto_limpio(valor):
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return ""
    texto = str(valor).strip()
    return texto[:-2] if texto.endswith(".0") else texto


def es_codigo_provisional(codigo):
    c = texto_normalizado(codigo)
    return c in {"", "-"} or any(
        t in c for t in ("PENDIENTE ASIGNAR", "PENDIENTE DE ASIGNAR", "POR ASIGNAR")
    )


def nombre_de_plan(alcance):
    """Nombre canónico del plan; si no se reconoce, se deja el texto original."""
    n = texto_normalizado(alcance)
    if "INYECCION" in n:
        return "PUNTOS DE INYECCIÓN"
    if "DEAD" in n:
        return "DEAD LEG"
    if "SULFID" in n:
        return "SULFIDACIÓN"
    return texto_limpio(alcance)


def normalizar_planes(df_entrada):
    df = df_entrada.copy()
    for columna in COLUMNAS:
        if columna not in df.columns:
            df[columna] = ""
    df = df[COLUMNAS].fillna("")
    for columna in COLUMNAS:
        df[columna] = df[columna].map(texto_limpio)
    df["MES"] = df["MES"].map(lambda v: v.upper())
    df["ALCANCE DEL SERVICIO"] = df["ALCANCE DEL SERVICIO"].map(nombre_de_plan)
    return df.reset_index(drop=True)


def es_retirada(fila):
    return (
        "RETIRADO" in texto_normalizado(fila[COL_ESTADO])
        or "RETIRADO" in texto_normalizado(fila["OBSERVACIÓN"])
        or texto_normalizado(fila["VALORIZACIÓN"]) == "RETIRADO"
    )


def clave_informe(fila):
    mes = texto_limpio(fila["MES"])
    if es_codigo_provisional(fila["CODIGO DE INFORME"]):
        return f"{mes}|SIN-CODIGO-GRUPO|{texto_normalizado(fila['GRUPO DE TUBERÍAS'])}"
    return f"{mes}|{texto_limpio(fila['CODIGO DE INFORME'])}"


def _estado_de_informe(estados):
    normalizados = [texto_normalizado(e) for e in estados]
    if any("PENDIENTE INSPECCION" in e for e in normalizados):
        return PEND_INSPECCION
    if any("PENDIENTE ELABORACION" in e for e in normalizados):
        return PEND_ELABORACION
    if all("FINALIZADO" in e or "100%" in e for e in normalizados):
        return FINALIZADO
    return OTRO


def _primero_no_vacio(serie):
    for valor in serie:
        if valor:
            return valor
    return ""


def _valorizacion(serie):
    valores = [texto_normalizado(v) for v in serie]
    if valores and all(v == "SI" for v in valores):
        return "SI"
    return _primero_no_vacio(serie)


def es_revision_fiabilidad(observacion):
    o = texto_normalizado(observacion)
    return "FIABILIDAD" in o and "REVISION" in o


def es_pendiente_especialista(observacion):
    return "PENDIENTE REVISION POR EL ESPECIALISTA" in texto_normalizado(observacion)


def es_revisado_especialista(observacion):
    o = texto_normalizado(observacion)
    return (
        "REV. POR EL ESPECIALISTA" in o
        or "REVISION POR EL ESPECIALISTA" in o
        or "REVISADO POR ESPECIALISTA" in o
    ) and "PENDIENTE" not in o


ETAPA_SIN_ENVIAR = "Por enviar a revisión"
ETAPA_PEND_ESP = "Pend. revisión especialista"
ETAPA_REVISADO = "Revisado por especialista"
ETAPA_FIABILIDAD = "En revisión fiabilidad"
ETAPA_VALORIZADO = "Valorizado"
OBS_PEND_ESP = "Pendiente revisión por el especialista - Ademinsac"
OBS_REVISADO = "INFORME REVISADO POR ESPECIALISTA"


def etapa_de_informe(estado, observacion, valorizacion):
    """Etapa de control del informe: elaboración, revisión del especialista, valorización."""
    if estado != FINALIZADO:
        return estado
    if texto_normalizado(valorizacion) == "SI":
        return ETAPA_VALORIZADO
    if es_revisado_especialista(observacion):
        return ETAPA_REVISADO
    if es_pendiente_especialista(observacion):
        return ETAPA_PEND_ESP
    if es_revision_fiabilidad(observacion):
        return ETAPA_FIABILIDAD
    return ETAPA_SIN_ENVIAR


def aplicar_accion(df, clave, accion, especialista=""):
    """Actualiza las filas activas de un informe y devuelve la matriz nueva.

    accion: "enviar_revision" | "revisado" | "valorizar" | "quitar_valorizacion".
    """
    df = normalizar_planes(df)
    mascara = df.apply(
        lambda f: not es_retirada(f) and clave_informe(f) == clave, axis=1
    )
    sufijo = f" (Especialista: {especialista})" if especialista else ""
    if accion == "enviar_revision":
        df.loc[mascara, "OBSERVACIÓN"] = OBS_PEND_ESP + sufijo
        df.loc[mascara, "VALORIZACIÓN"] = "Pendiente"
    elif accion == "revisado":
        df.loc[mascara, "OBSERVACIÓN"] = OBS_REVISADO + sufijo
    elif accion == "valorizar":
        df.loc[mascara, "VALORIZACIÓN"] = "SI"
        df.loc[mascara, "OBSERVACIÓN"] = ""
    elif accion == "quitar_valorizacion":
        # Solo se valoriza un informe ya revisado: vuelve a ese estado.
        df.loc[mascara, "VALORIZACIÓN"] = "Pendiente"
        df.loc[mascara, "OBSERVACIÓN"] = OBS_REVISADO
    else:
        raise ValueError(accion)
    return df


def procesar_planes(df_planes):
    """Devuelve activos, retiradas, informes (una fila por informe) y KPIs."""
    df = normalizar_planes(df_planes)
    mascara = df.apply(es_retirada, axis=1) if len(df) else pd.Series(dtype=bool)
    retiradas = df[mascara].copy()
    activos = df[~mascara].copy()
    activos["CLAVE"] = activos.apply(clave_informe, axis=1) if len(activos) else ""

    filas = []
    for clave, g in activos.groupby("CLAVE", sort=False):
        filas.append({
            "CLAVE": clave,
            "PLAN": g["ALCANCE DEL SERVICIO"].iloc[0],
            "MES": g["MES"].iloc[0],
            "UNIDAD": ", ".join(sorted({u for u in g["UNIDAD"] if u})),
            "CODIGO DE INFORME": g["CODIGO DE INFORME"].iloc[0],
            "GRUPO DE TUBERÍAS": g["GRUPO DE TUBERÍAS"].iloc[0],
            "PUNTOS": len(g),
            "ESTADO": _estado_de_informe(g[COL_ESTADO]),
            "RESPONSABLE": _primero_no_vacio(g["RESPONSABLE"]),
            "OBSERVACIÓN": _primero_no_vacio(g["OBSERVACIÓN"]),
            "VALORIZACIÓN": _valorizacion(g["VALORIZACIÓN"]),
        })
    for f in filas:
        f["ETAPA"] = etapa_de_informe(f["ESTADO"], f["OBSERVACIÓN"], f["VALORIZACIÓN"])
    informes = pd.DataFrame(
        filas,
        columns=["CLAVE", "PLAN", "MES", "UNIDAD", "CODIGO DE INFORME",
                 "GRUPO DE TUBERÍAS", "PUNTOS", "ESTADO", "RESPONSABLE",
                 "OBSERVACIÓN", "VALORIZACIÓN", "ETAPA"],
    )
    if not informes.empty:
        informes["_m"] = informes["MES"].map(
            lambda m: ORDEN_MESES.index(m) if m in ORDEN_MESES else 99
        )
        informes["_p"] = informes["PLAN"].map(
            lambda p: PLANES.index(p) if p in PLANES else 99
        )
        informes = (
            informes.sort_values(["_p", "_m"], kind="stable")
            .drop(columns=["_m", "_p"])
            .reset_index(drop=True)
        )

    return {
        "activos": activos,
        "retiradas": retiradas,
        "informes": informes,
        "kpis": calcular_kpis(informes, retiradas),
    }


def _kpis_de(informes, retiradas):
    def n(df):
        return int(len(df))

    def pts(df):
        return int(df["PUNTOS"].sum())

    def etapa(*nombres):
        return informes[informes["ETAPA"].isin(nombres)]

    fin = informes[informes["ESTADO"] == FINALIZADO]
    p_ins = informes[informes["ESTADO"] == PEND_INSPECCION]
    p_ela = informes[informes["ESTADO"] == PEND_ELABORACION]
    val = etapa(ETAPA_VALORIZADO)
    sin_val = fin[fin["ETAPA"] != ETAPA_VALORIZADO]
    sin_enviar = etapa(ETAPA_SIN_ENVIAR)
    pend_esp = etapa(ETAPA_PEND_ESP)
    revisados = etapa(ETAPA_REVISADO)
    fiab = etapa(ETAPA_FIABILIDAD)
    total = n(informes)
    return {
        "informes": total,
        "puntos": pts(informes),
        "finalizados": n(fin),
        "puntos_finalizados": pts(fin),
        "pend_elaborar": total - n(fin),  # todo lo que aún no está finalizado
        "puntos_pend_elaborar": pts(informes) - pts(fin),
        "pend_inspeccion": n(p_ins),
        "puntos_pend_inspeccion": pts(p_ins),
        "pend_elaboracion": n(p_ela),
        "puntos_pend_elaboracion": pts(p_ela),
        "retirados": int(len(retiradas)),
        "valorizados": n(val),
        "puntos_valorizados": pts(val),
        "pend_valorizar": n(sin_val),
        "puntos_pend_valorizar": pts(sin_val),
        "sin_enviar": n(sin_enviar),
        "puntos_sin_enviar": pts(sin_enviar),
        "por_revisar": n(pend_esp),
        "puntos_por_revisar": pts(pend_esp),
        # Los valorizados ya pasaron por la revisión del especialista.
        "revisados": n(revisados),
        "puntos_revisados": pts(revisados),
        "en_revision": n(fiab),
        "puntos_en_revision": pts(fiab),
        "avance": (n(fin) / total * 100) if total else 0.0,
        "avance_valorizacion": (n(val) / total * 100) if total else 0.0,
    }


def calcular_kpis(informes, retiradas):
    """KPIs del total y de cada plan (claves: "TOTAL" y el nombre del plan)."""
    kpis = {"TOTAL": _kpis_de(informes, retiradas)}
    planes = list(PLANES) + sorted(
        p for p in set(informes["PLAN"]) | set(retiradas["ALCANCE DEL SERVICIO"])
        if p and p not in PLANES
    )
    for plan in planes:
        kpis[plan] = _kpis_de(
            informes[informes["PLAN"] == plan],
            retiradas[retiradas["ALCANCE DEL SERVICIO"] == plan],
        )
    return kpis


def resumen_por_plan(kpis):
    """Tabla resumen: una fila por plan y fila TOTAL."""
    filas = []
    for plan, k in kpis.items():
        if plan == "TOTAL":
            continue
        filas.append({
            "PLAN": plan,
            "INFORMES": k["informes"],
            "PUNTOS": k["puntos"],
            "FINALIZADOS": k["finalizados"],
            "PEND. INSPECCIÓN": k["pend_inspeccion"],
            "PEND. ELABORACIÓN": k["pend_elaboracion"],
            "RETIRADOS": k["retirados"],
            "REVISADOS": k["revisados"],
            "VALORIZADOS": k["valorizados"],
            "% AVANCE": f"{k['avance']:.1f}%",
        })
    t = kpis["TOTAL"]
    filas.append({
        "PLAN": "TOTAL",
        "INFORMES": t["informes"],
        "PUNTOS": t["puntos"],
        "FINALIZADOS": t["finalizados"],
        "PEND. INSPECCIÓN": t["pend_inspeccion"],
        "PEND. ELABORACIÓN": t["pend_elaboracion"],
        "RETIRADOS": t["retirados"],
        "REVISADOS": t["revisados"],
        "VALORIZADOS": t["valorizados"],
        "% AVANCE": f"{t['avance']:.1f}%",
    })
    tabla = pd.DataFrame(filas)
    tabla.index = range(1, len(tabla) + 1)
    return tabla


def resumen_por_mes(informes):
    """Informes por mes (filas) y plan (columnas), con fila TOTAL."""
    if informes.empty:
        return pd.DataFrame()
    cuenta = informes.pivot_table(
        index="MES", columns="PLAN", values="CLAVE", aggfunc="nunique", fill_value=0
    )
    columnas = [p for p in PLANES if p in cuenta.columns] + [
        p for p in cuenta.columns if p not in PLANES
    ]
    cuenta = cuenta[columnas]
    meses = [m for m in ORDEN_MESES if m in cuenta.index] + [
        m for m in cuenta.index if m not in ORDEN_MESES
    ]
    cuenta = cuenta.loc[meses]
    cuenta["TOTAL"] = cuenta.sum(axis=1)
    cuenta.loc["TOTAL"] = cuenta.sum()
    cuenta = cuenta.astype(int).reset_index().rename(columns={"MES": "MES"})
    cuenta.columns.name = None
    cuenta.index = range(1, len(cuenta) + 1)
    return cuenta


# ---------------------------------------------------------------- Persistencia y Excel
def cargar_json(ruta):
    if not os.path.exists(ruta):
        return pd.DataFrame(columns=COLUMNAS)
    try:
        with open(ruta, "r", encoding="utf-8") as archivo:
            return normalizar_planes(pd.DataFrame(json.load(archivo)))
    except Exception:
        return pd.DataFrame(columns=COLUMNAS)


def guardar_json(df, ruta):
    normalizar_planes(df).to_json(ruta, orient="records", force_ascii=False)


def leer_excel_planes(archivo):
    """Lee la hoja CONTROL (o la primera) de un Excel de planes y lo normaliza."""
    libro = pd.ExcelFile(archivo)
    hoja = "CONTROL" if "CONTROL" in libro.sheet_names else libro.sheet_names[0]
    df = pd.read_excel(libro, sheet_name=hoja, dtype=str)
    mapa = {str(c).strip().upper(): c for c in df.columns}
    renombre = {
        mapa[c.strip().upper()]: c for c in COLUMNAS if c.strip().upper() in mapa
    }
    return normalizar_planes(df.rename(columns=renombre))


def excel_con_formato(df, nombre_hoja="CONTROL"):
    """Excel con el mismo formato que el respaldo de la matriz general."""
    salida = io.BytesIO()
    datos = df.copy().fillna("")
    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        datos.to_excel(escritor, index=False, sheet_name=nombre_hoja[:31])
        hoja = escritor.book[nombre_hoja[:31]]
        hoja.freeze_panes = "A2"
        hoja.row_dimensions[1].height = 30
        relleno = PatternFill("solid", fgColor="0E2A47")
        for celda in hoja[1]:
            celda.font = Font(color="FFFFFF", bold=True)
            celda.fill = relleno
            celda.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for columna in hoja.columns:
            letra = get_column_letter(columna[0].column)
            ancho = max(len(texto_limpio(c.value)) for c in columna)
            hoja.column_dimensions[letra].width = min(max(ancho + 2, 12), 42)
            for celda in columna[1:]:
                celda.alignment = Alignment(vertical="top", wrap_text=True)
        if len(datos.columns) and len(datos):
            tabla = Table(
                displayName=f"Tabla_{datetime.now():%H%M%S%f}",
                ref=f"A1:{get_column_letter(len(datos.columns))}{len(datos) + 1}",
            )
            tabla.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True
            )
            hoja.add_table(tabla)
    salida.seek(0)
    return salida.getvalue()
