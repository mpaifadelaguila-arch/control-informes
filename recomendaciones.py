"""
recomendaciones.py — motor de reglas (SIN IA) para redactar "Hallazgo" y
"Recomendación Técnica Aplicable" a partir de datos estructurados que el
inspector de campo llena en el VT-CHECK LIST, siguiendo el catálogo de casos
típicos de COMPLEMENTO/COMPENDIO TÉCNICO UNIFICADO DE HALLAZGOS Y
RECOMENDACIONES TÉCNICAS.REV.1.docx y las reglas críticas de
COMPLEMENTO/ROL_Y_OBJETIVO.REV2.txt.

No hay llamadas a ningún modelo de lenguaje: cada fila del checklist trae un
código de "CASO" (elegido de una lista cerrada) más un puñado de variables de
campo (NPS, longitud, cantidad, ubicación, color, etc.); este módulo busca el
caso en el catálogo, valida/aplica las reglas críticas (severidad leve
puntual -> sin recomendación, tramo no accesible -> andamios, etc.) y rellena
la plantilla oficial con los valores recibidos.

El catálogo de casos (CATALOGO) y las reglas de sugerencia automática
(REGLAS_SUGERENCIA) YA NO están escritos a mano en este archivo: se cargan en
tiempo de ejecución desde Catalogo_Hallazgos_Recomendaciones.xlsx (raíz del
repositorio), igual que BASE_DE_DATOS_DE_LINEAS_FASE1.xlsx -- se puede
agregar o corregir un caso editando ese Excel, sin tocar código ni volver a
desplegar la app (ver cargar_catalogo_desde_excel() y recargar_catalogo()
más abajo, y la hoja "Léame" del propio Excel).

Uso típico (desde checklist.py):

    from recomendaciones import generar_hallazgo_y_recomendacion
    resultado = generar_hallazgo_y_recomendacion(fila_dict)
    if resultado is not None:
        hallazgo, recomendacion = resultado.hallazgo, resultado.recomendacion
"""
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional


SIN_DATO = "SIN DATO"

# Casos en los que, aunque el tramo esté en altura / no accesible, el
# hallazgo SÍ se describe y recomienda de forma directa (excepción de la
# sección 3 de ROL_Y_OBJETIVO.REV2.txt: fuga en bridas, fuga en válvulas,
# pandeo/deformación de tubería).
CASOS_VISIBLES_PESE_A_ALTURA = {
    "BRIDA_FUGA_EMPAQUE",
    "VALVULA_CONTROL_FUGA_CUERPO_BONETE",
    "TUBERIA_PANDEO_DEFORMACION",
}

# Severidades que, según la matriz de decisión (regla 1 de ROL_Y_OBJETIVO),
# NUNCA generan recomendación técnica: falla de pintura o corrosión leve y
# puntual, el componente conserva su integridad.
SEVERIDADES_SIN_RECOMENDACION = {"LEVE PUNTUAL", "LEVE LOCALIZADA"}

CASO_ANDAMIO = "TUBERIA_TRAMO_NO_ACCESIBLE"


def _fmt(valor, default=None):
    """Devuelve el valor limpio como string, o SIN_DATO/placeholder si falta."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return default if default is not None else SIN_DATO
    return str(valor).strip()


@dataclass
class Caso:
    id: str
    categoria: str
    etiqueta: str
    hallazgo_tpl: str
    recomendacion_tpl: str
    variables: tuple = field(default_factory=tuple)
    defaults: dict = field(default_factory=dict)


@dataclass
class Resultado:
    caso_id: str
    hallazgo: str
    recomendacion: str


def _v(datos, clave, defecto=SIN_DATO):
    return _fmt(datos.get(clave), defecto)


# Cuando el checklist no trae un dato de diámetro/longitud/material/schedule,
# en vez de insertar "SIN DATO" en la frase final (lo que se ve poco
# profesional y confunde, a pedido del usuario) se omite esa cláusula
# completa: la recomendación se redacta obviando ese dato, en vez de
# mostrarlo como faltante.
_RE_OMITIR_LONGITUD_PARENTESIS = re.compile(
    r"\s*\(longitud aprox\.\s+SIN DATO metros\)", re.IGNORECASE
)
_RE_OMITIR_LONGITUD_BARE = re.compile(
    r"\s*(?:de\s+)?longitud aprox\.\s+SIN DATO metros\b", re.IGNORECASE
)
_RE_OMITIR_APROXIMADAMENTE = re.compile(
    r"\s*(?:de\s+)?aproximadamente\s+SIN DATO metros\b", re.IGNORECASE
)
_RE_OMITIR_MEDIDA = re.compile(
    r"\s*(?:de\s+)?(?:NPS|Sch|material|tipo)\s+SIN DATO\b", re.IGNORECASE
)
# Caso "de {tipo}" sin la palabra "tipo" delante (p.ej. "válvula de {tipo}
# NPS..." -> "válvula de SIN DATO NPS..."): el conector "de" va pegado
# directo a SIN DATO, sin otra etiqueta entre medio.
_RE_OMITIR_DE_BARE = re.compile(r"\s*\bde\s+SIN DATO\b", re.IGNORECASE)
# Caso "{tag} / NPS {nps}" (válvulas de control): si falta el TAG, la barra
# queda pegada a SIN DATO por el lado izquierdo o derecho según qué dato
# falte; se resuelve cada combinación por separado.
_RE_OMITIR_SLASH_AMBOS = re.compile(r"\bSIN DATO\s*/\s*SIN DATO\b", re.IGNORECASE)
_RE_OMITIR_SLASH_IZQ = re.compile(r"\bSIN DATO\s*/\s*", re.IGNORECASE)
_RE_OMITIR_SLASH_DER = re.compile(r"\s*/\s*SIN DATO\b", re.IGNORECASE)
# Red de seguridad final: pase lo que pase con la variable (un caso nuevo,
# una combinación no prevista por las reglas específicas de arriba), esto
# GARANTIZA que la palabra "SIN DATO" nunca llegue al texto entregado --
# a pedido explícito y estricto del usuario. Primero intenta quitar
# también la preposición/paréntesis que la introduce (para que el texto
# quede lo más natural posible); si ni así calza con un patrón conocido,
# el último recurso es borrar únicamente las dos palabras.
_RE_OMITIR_PREPOSICION_GENERICA = re.compile(
    r"\s*\((?:[^()]*\bSIN DATO\b[^()]*)\)", re.IGNORECASE
)
_RE_OMITIR_PREP_SIN_DATO = re.compile(
    r"\s*\b(?:en|de|del|de la|de el|con|desde|hacia|sobre|por|a)\s+SIN DATO\b",
    re.IGNORECASE,
)
_RE_SIN_DATO_RESIDUAL = re.compile(r"SIN DATO", re.IGNORECASE)


def _omitir_datos_faltantes(texto):
    """Limpia de una frase ya redactada cualquier cláusula (diámetro NPS,
    longitud, schedule, material, tipo de válvula/soporte...) que haya
    quedado en SIN_DATO por no venir en el comentario de campo, dejando la
    oración gramaticalmente correcta en vez de mostrar el literal
    "SIN DATO"."""
    if not texto:
        return texto
    t = _RE_OMITIR_LONGITUD_PARENTESIS.sub("", texto)
    t = _RE_OMITIR_LONGITUD_BARE.sub("", t)
    t = _RE_OMITIR_APROXIMADAMENTE.sub("", t)
    t = _RE_OMITIR_MEDIDA.sub("", t)
    t = _RE_OMITIR_DE_BARE.sub("", t)
    t = _RE_OMITIR_SLASH_AMBOS.sub("", t)
    t = _RE_OMITIR_SLASH_IZQ.sub("", t)
    t = _RE_OMITIR_SLASH_DER.sub("", t)
    # Red de seguridad: cualquier "SIN DATO" que ninguna regla específica
    # de arriba haya anticipado se quita igual, primero con su paréntesis
    # o preposición si los tiene, y si no, a secas -- nunca debe sobrevivir.
    t = _RE_OMITIR_PREPOSICION_GENERICA.sub("", t)
    t = _RE_OMITIR_PREP_SIN_DATO.sub("", t)
    t = _RE_SIN_DATO_RESIDUAL.sub("", t)
    t = _limpiar_puntuacion(t)
    return t.strip()


def _limpiar_puntuacion(t):
    t = re.sub(r"\s+([.,;:])", r"\1", t)  # sin espacio antes de puntuación
    t = re.sub(r"(?:,\s*){2,}", ", ", t)  # dos o más comas seguidas (p.ej. al
    # perderse NPS, Sch y material a la vez) colapsadas en una sola
    t = re.sub(r",\s*\.", ".", t)  # coma justo antes del punto final
    t = re.sub(r"\(\s*\)", "", t)  # paréntesis vacíos
    t = re.sub(r"\s{2,}", " ", t)  # espacios dobles
    t = re.sub(r"^,\s*", "", t)  # coma colgando al inicio de la frase
    return t


# Varias plantillas del catálogo escriben "válvula(s)", "unión(es)
# bridada(s)", etc. para servir tanto al singular como al plural sin
# necesitar dos redacciones distintas. Si la cantidad real resulta ser 1,
# no tiene sentido mostrar el "(s)"/"(es)" -- a pedido del usuario, se
# resuelve automáticamente a la forma singular o plural real según el
# número que haya quedado justo antes del bloque de palabra(s).
# La cantidad puede ser un número real ("3") o, cuando el comentario de
# campo no la especifica, el valor por defecto "varias"/"varios" (nunca
# singular -- si no se sabe cuántas son, no se afirma que sea solo una).
_RE_CANTIDAD_PLURAL = re.compile(
    r"(\d+|varias?|varios?)\s+((?:\S*?\((?:es|s)\)\s*)+)", re.IGNORECASE
)
_RE_TOKEN_PLURAL = re.compile(r"(\S*?)\((es|s)\)")


def _resolver_plurales(texto):
    if not texto:
        return texto

    def _procesar_bloque(m):
        numero, bloque = m.group(1), m.group(2)
        es_singular = numero.isdigit() and int(numero) == 1

        def _token(tm):
            palabra, sufijo = tm.group(1), tm.group(2)
            if es_singular:
                return palabra
            # "unión(es)" -> "uniones", no "uniónes": al pluralizar una
            # palabra terminada en "-ión" se pierde la tilde (regla
            # ortográfica del español), no se le pega "es" tal cual.
            if sufijo == "es" and palabra.lower().endswith("ión"):
                return palabra[:-3] + "iones"
            return f"{palabra}{sufijo}"

        # Sin .strip(): el espacio final capturado en `bloque` (antes de la
        # palabra siguiente, p.ej. "válvula(s) de...") debe conservarse tal
        # cual, o quedarían dos palabras pegadas ("válvulade").
        return f"{numero} {_RE_TOKEN_PLURAL.sub(_token, bloque)}"

    return _RE_CANTIDAD_PLURAL.sub(_procesar_bloque, texto)


_VERBOS_IMPERATIVOS = (
    "Realizar", "Efectuar", "Instalar", "Reemplazar", "Aplicar", "Retirar",
    "Reparar", "Limpiar", "Corregir", "Reforzar", "Verificar", "Solicitar",
    "Programar", "Colocar", "Restablecer", "Adecuar", "Ejecutar",
)
_RE_VERBO_INICIAL = re.compile(
    r"^(?:" + "|".join(_VERBOS_IMPERATIVOS) + r")\s+", re.IGNORECASE
)


# Varios casos del catálogo ya usan "así como" dentro de su propia
# redacción (p.ej. "...así como su lubricación..."). Si la recomendación
# combinada repitiera siempre el mismo conector, quedaría "...así como su
# lubricación... Así como el reemplazo..." -- repetitivo y confuso. Se
# elige el primer conector de esta lista que NO aparezca ya en el texto
# acumulado hasta ese punto.
_CONECTORES_CONTINUACION = ("Así como", "Asimismo,", "Adicionalmente,")


def _elegir_conector(texto_previo):
    previo_norm = texto_previo.lower()
    for conector in _CONECTORES_CONTINUACION:
        if conector.rstrip(",").lower() not in previo_norm:
            return conector
    return _CONECTORES_CONTINUACION[-1]


def _como_continuacion(texto, conector="Así como"):
    """Convierte la recomendación de un caso ADICIONAL (el 2do, 3er... caso
    detectado en la misma observación) en la continuación de una sola
    oración, en vez de una nueva oración imperativa aparte: se le quita el
    verbo inicial ("Realizar", "Efectuar"...) y se antepone el conector
    elegido -- nunca un encabezado "Recomendación:" ni viñetas, que hacían
    la recomendación combinada más larga y menos legible de lo necesario."""
    t = texto.strip()
    if t.lower().startswith("recomendación:"):
        t = t.split(":", 1)[1].strip()
    t = t.lstrip("•").strip()
    t = _RE_VERBO_INICIAL.sub("", t)
    if t:
        t = t[0].lower() + t[1:]
    return f"{conector} {t}"


# ==============================================================================
# CATÁLOGO — cargado desde Catalogo_Hallazgos_Recomendaciones.xlsx (hoja
# "Casos"), extraído originalmente 1:1 del COMPENDIO REV.1 (secciones A-H).
# Ver cargar_catalogo_desde_excel() más abajo para el formato exacto.
# ==============================================================================
RUTA_CATALOGO_DEFAULT = Path(__file__).resolve().parent / "Catalogo_Hallazgos_Recomendaciones.xlsx"


def _parsear_variables(valor):
    return tuple(v.strip() for v in (valor or "").split(",") if v.strip())


def _parsear_defaults(valor):
    defaults = {}
    for par in (valor or "").split(";"):
        if "=" in par:
            k, v = par.split("=", 1)
            k = k.strip()
            if k:
                defaults[k] = v.strip()
    return defaults


def _parsear_categorias(valor):
    t = (valor or "").strip()
    if not t or t.upper() == "CUALQUIERA":
        return (None,)
    return tuple(c.strip() for c in t.split(";") if c.strip())


def _parsear_palabras(valor):
    return tuple(p.strip() for p in (valor or "").split("|") if p.strip())


def cargar_catalogo_desde_excel(ruta):
    """Lee el catálogo de casos y las reglas de sugerencia automática desde
    el Excel editable (hojas "Casos" y "Reglas_Sugerencia" -- formato
    documentado en la hoja "Léame" de ese mismo archivo). Devuelve
    (catalogo_dict, reglas_lista) en el mismo formato que antes vivía
    escrito a mano en este módulo, para que el resto del motor no note la
    diferencia."""
    import openpyxl

    wb = openpyxl.load_workbook(ruta, data_only=True)

    catalogo = {}
    for row in wb["Casos"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        cid, categoria, etiqueta, hallazgo_tpl, recomendacion_tpl, variables_str, defaults_str = (
            list(row) + [None] * 7
        )[:7]
        cid = str(cid).strip()
        catalogo[cid] = Caso(
            id=cid,
            categoria=(categoria or "").strip(),
            etiqueta=(etiqueta or "").strip(),
            hallazgo_tpl=hallazgo_tpl or "",
            recomendacion_tpl=recomendacion_tpl or "",
            variables=_parsear_variables(variables_str),
            defaults=_parsear_defaults(defaults_str),
        )

    reglas_por_num = {}
    for row in wb["Reglas_Sugerencia"].iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        regla_num, caso_id, categorias_str, grupo_num, palabras_str, excluidas_str = (
            list(row) + [None] * 6
        )[:6]
        entrada = reglas_por_num.setdefault(int(regla_num), {
            "caso_id": str(caso_id).strip(),
            "categorias_str": categorias_str,
            "grupos": {},
            "excluidas_str": None,
        })
        entrada["grupos"][int(grupo_num)] = _parsear_palabras(palabras_str)
        if excluidas_str:
            entrada["excluidas_str"] = excluidas_str

    reglas = []
    for num in sorted(reglas_por_num):
        e = reglas_por_num[num]
        categorias = _parsear_categorias(e["categorias_str"])
        grupos = [e["grupos"][g] for g in sorted(e["grupos"])]
        excluidas = _parsear_palabras(e["excluidas_str"])
        if excluidas:
            reglas.append((e["caso_id"], categorias, grupos, excluidas))
        else:
            reglas.append((e["caso_id"], categorias, grupos))

    return catalogo, reglas


def _cargar_catalogo_inicial():
    try:
        return cargar_catalogo_desde_excel(RUTA_CATALOGO_DEFAULT)
    except Exception as e:
        print(
            f"[!] No se pudo cargar el catálogo de hallazgos y recomendaciones "
            f"desde {RUTA_CATALOGO_DEFAULT}: {e}. El motor seguirá funcionando, "
            f"pero ningún hallazgo tendrá recomendación automática (quedarán "
            f"todos PENDIENTES de redacción manual) hasta que el archivo esté "
            f"disponible."
        )
        return {}, []


CATALOGO, REGLAS_SUGERENCIA = _cargar_catalogo_inicial()


def recargar_catalogo(ruta=None):
    """Vuelve a leer el catálogo desde el Excel y actualiza CATALOGO y
    REGLAS_SUGERENCIA en caliente -- para que una edición de ese archivo
    surta efecto sin reiniciar la app, igual que
    BASE_DE_DATOS_DE_LINEAS_FASE1.xlsx. Se llama automáticamente al inicio
    de cada generación de informe (ver pages/1_Elaboracion_de_Informe.py)."""
    global CATALOGO, REGLAS_SUGERENCIA
    CATALOGO, REGLAS_SUGERENCIA = cargar_catalogo_desde_excel(ruta or RUTA_CATALOGO_DEFAULT)

CATEGORIAS = sorted({c.categoria for c in CATALOGO.values()})


def casos_por_categoria(categoria):
    return [c for c in CATALOGO.values() if c.categoria == categoria.strip().upper()]


def _severidad_norm(v):
    return _fmt(v, "").strip().upper()


def generar_hallazgo_y_recomendacion(datos: dict) -> Optional[Resultado]:
    """Recibe un dict con, al menos, la clave "caso" (id del catálogo o su
    etiqueta) y las variables necesarias (nps, longitud, cantidad, ubicacion,
    tipo, tag, color_actual, color_norma, elemento, material, schedule,
    severidad) y devuelve un Resultado(hallazgo, recomendacion) ya redactado,
    o None si la regla crítica de severidad determina que NO corresponde
    recomendación (leve puntual / localizada -> aceptable, se omite).

    `datos` también puede traer "accesible" ("SI"/"NO") y el motor
    sustituye automáticamente el caso por TUBERIA_TRAMO_NO_ACCESIBLE cuando
    accesible="NO", salvo que el caso elegido esté en
    CASOS_VISIBLES_PESE_A_ALTURA (fuga de brida/válvula, pandeo)."""
    caso_id = _fmt(datos.get("caso"), "").strip()
    if not caso_id:
        return None

    caso = CATALOGO.get(caso_id.upper())
    if caso is None:
        # Tolerar que venga la etiqueta en español en vez del id.
        caso_id_norm = caso_id.strip().lower()
        for c in CATALOGO.values():
            if c.etiqueta.strip().lower() == caso_id_norm:
                caso = c
                break
    if caso is None:
        raise ValueError(
            f"Caso de hallazgo no reconocido en el catálogo: {caso_id!r}. "
            "Revisa la columna CASO del VT-CHECK LIST."
        )

    severidad = _severidad_norm(datos.get("severidad"))
    if severidad in SEVERIDADES_SIN_RECOMENDACION:
        return None

    accesible = _fmt(datos.get("accesible"), "SI").strip().upper()
    if accesible in ("NO", "N") and caso.id not in CASOS_VISIBLES_PESE_A_ALTURA:
        caso = CATALOGO[CASO_ANDAMIO]

    valores = dict(caso.defaults)
    for var in caso.variables:
        # Si el checklist no trae el dato, se respeta el valor por defecto
        # del propio caso (si lo tiene) en vez de caer directo a "SIN DATO".
        valores[var] = _v(datos, var, valores.get(var, SIN_DATO))
    # Variables genéricas siempre disponibles aunque el caso no las declare
    # explícitamente en `variables` (por si el usuario las incluye en el
    # texto manualmente vía formato libre no usado aquí).
    if "severidad" in caso.hallazgo_tpl and "severidad" not in valores:
        valores["severidad"] = _v(datos, "severidad", caso.defaults.get("severidad", SIN_DATO))

    try:
        hallazgo = caso.hallazgo_tpl.format(**valores)
        recomendacion = caso.recomendacion_tpl.format(**valores)
    except KeyError as e:
        raise ValueError(
            f"Falta la variable {e} en el checklist para el caso {caso.id!r} "
            f"({caso.etiqueta})."
        )

    hallazgo = _resolver_plurales(_omitir_datos_faltantes(hallazgo))
    recomendacion = _resolver_plurales(_omitir_datos_faltantes(recomendacion))
    return Resultado(caso_id=caso.id, hallazgo=hallazgo, recomendacion=recomendacion)


# ==============================================================================
# SUGERENCIA DE CASO A PARTIR DE TEXTO LIBRE (formato real del VT-CHECK LIST)
# ==============================================================================
# El VT-CHECK LIST estándar de campo (formato fijo AD-UN05-TL-TUB-VT) NO trae
# una columna "CASO": el inspector escribe el hallazgo como texto libre en la
# columna "Comentario", por categoría (Recubrimientos, Bridas, Válvulas,
# Soportes, etc.) y marca A/O/R/NA. Esta sección NO usa ningún modelo de
# lenguaje: aplica un conjunto pequeño y explícito de reglas de palabras
# clave (categoría + condición) para sugerir, con alta confianza, cuál caso
# del CATALOGO corresponde, y solo cuando hay evidencia clara. Si ninguna
# regla calza con confianza, se devuelve None y el hallazgo queda marcado
# para redacción manual del especialista (nunca se inventa una
# recomendación técnica sobre una base ambigua).

RE_NPS = re.compile(r'(\d+(?:\s+\d/\d)?"|\d/\d"?)')
# El patrón genérico de arriba encuentra CUALQUIER número-comillas, sin
# importar a qué se refiere -- en un comentario real como "01 de 8
# espárrago de 5/8" en la brida de NPS 4"..." capturaba el diámetro del
# espárrago (5/8") en vez del NPS real de la línea/brida (4"), porque
# aparece primero en el texto. Se prueba primero, explícitamente, el
# número que viene inmediatamente después de la palabra "NPS" (o "Ø"),
# y solo si no aparece se cae al patrón genérico de respaldo.
RE_NPS_ETIQUETADO = re.compile(r'(?:NPS|Ø)\s*(\d+\s+\d/\d"?|\d/\d"?|\d+"?)', re.IGNORECASE)
# Dos formas reales de escribir la cantidad en el checklist: entre
# paréntesis "(3)" (grupo 1), o un número suelto -- a veces con cero a la
# izquierda, "01 unión bridada" -- justo antes de uno de los sustantivos
# contables típicos del catálogo (grupo 2). Nunca un número cualquiera
# (evita capturar el NPS o una fecha): exige que el sustantivo aparezca
# pegado al número.
RE_CANTIDAD = re.compile(
    r"\((\d{1,3})\)"
    r"|\b0*(\d{1,3})\s+(?:uni[oó]n(?:es)?|v[aá]lvulas?|abrazaderas?|soportes?|"
    r"juntas?|esp[aá]rragos?|zonas?|sectores?|bridas?|tramos?|aberturas?|"
    r"conexi[oó]n(?:es)?)\b",
    re.IGNORECASE,
)
# Caso "N de M <sustantivo>" (p.ej. "01 de 8 espárrago"): el inspector da
# el TOTAL de elementos y, aparte, cuántos de esos están afectados -- la
# cantidad que importa para la recomendación es la AFECTADA (N), no el
# total (M). Sin este patrón, RE_CANTIDAD de arriba capturaba el número
# pegado al sustantivo (M, el total) en vez del real afectado (N). Se
# prueba primero, antes del patrón genérico.
RE_CANTIDAD_DE_TOTAL = re.compile(
    r"\b0*(\d{1,3})\s+de\s+0*(\d{1,3})\s+(?:esp[aá]rragos?|uni[oó]n(?:es)?|"
    r"v[aá]lvulas?|abrazaderas?|soportes?|juntas?|zonas?|sectores?|bridas?|tramos?)\b",
    re.IGNORECASE,
)
# El inspector a veces da varias longitudes juntas para varios sectores/
# tramos, antes de un único "metros" final (p.ej. "longitud aproximada de
# 1.5 y 0.5 metros"): el patrón anterior solo exigía un número PEGADO a
# "metros", así que se quedaba solo con el último (0.5) y perdía el
# primero (1.5) en silencio. Ahora se captura la secuencia completa de
# números separados por coma/"y" que preceden a "metros", tal cual la
# escribió el inspector.
RE_LONGITUD = re.compile(
    r"([\d]+(?:\.[\d]+)?(?:\s*(?:,|y)\s*[\d]+(?:\.[\d]+)?)*)\s*metros",
    re.IGNORECASE,
)
RE_TIPO_VALVULA = re.compile(
    r"v[aá]lvula[s]?\s+(?:de\s+|tipo\s+)?(compuerta|bola|globo|retenci[oó]n|check|"
    r"mariposa|aguja|tap[oó]n|diafragma|control|seguridad)",
    re.IGNORECASE,
)
RE_TIPO_SOPORTE = re.compile(r"\b(u-?bolt|spring hanger)\b", re.IGNORECASE)
RE_TRAMO_UBICACION = re.compile(r"tramo\s+(vertical|horizontal)", re.IGNORECASE)
# El inspector casi nunca escribe "tramo vertical/horizontal": lo habitual es
# ubicar el hallazgo con una frase locativa libre ("cerca de la válvula
# XV-101", "junto al soporte S-12", "a la altura del codo norte"...). Antes
# esto quedaba sin capturar y la plantilla mostraba "SIN DATO" aun teniendo
# el dato real presente en el comentario -- se agrega este patrón genérico
# como respaldo de RE_TRAMO_UBICACION.
RE_UBICACION_GENERICA = re.compile(
    r"(?:cerca de|cercan[oa]s?\s+a|pr[oó]xim[oa]s?\s+a|junto a|a la altura de|"
    r"en la zona de|en el [aá]rea de|a un costado de)(?:l)?\s+(.+?)(?:[.,;]|$)",
    re.IGNORECASE,
)

# Vocabulario de elementos/componentes/accesorios que el inspector suele
# nombrar en el comentario de campo. Se usa para que la Recomendación (y el
# Hallazgo, cuando el caso lo requiere) nombren EXPLÍCITAMENTE el elemento
# real inspeccionado -- nunca un texto genérico ("accesorios y
# componentes") -- cuando el propio comentario ya lo indica. Ordenado de
# más específico a más genérico para que, p.ej., "brida de orificio" se
# reconozca antes que la sola palabra "brida".
#
# NOTA: "protección ignífuga" queda fuera a propósito -- tiene su propio
# caso (AISLAMIENTO_PROTECCION_IGNIFUGA_AGRIETADA) donde {elemento} es lo
# que la protección ignífuga recubre (p.ej. una válvula o un soporte), no
# la protección en sí; incluirla aquí produciría un auto-referencia sin
# sentido ("protección ignífuga de protección ignífuga").
_ELEMENTOS_TIPICOS = (
    "brida de orificio", "placa orificio",
    "aislamiento térmico", "aislamiento termico",
    "base de concreto",
    "unión universal", "union universal",
    "cinta de valizamiento", "cinta señalizadora",
    "alambre de amarre",
    "brida", "válvula", "valvula", "abrazadera", "soporte",
    "venteo", "drenaje", "dren", "codo",
    "niples", "niple", "cabezal",
    "espárragos", "esparragos", "espárrago", "esparrago",
    "manómetro", "manometro", "indicador",
    "cinta", "alambre",
    "tubería", "tuberia",
)
RE_ELEMENTO = re.compile(
    "|".join(re.escape(e) for e in _ELEMENTOS_TIPICOS), re.IGNORECASE
)


def _contiene_alguna(texto, alternativas):
    return any(alt in texto for alt in alternativas)


def _extraer_elementos(texto):
    """Devuelve, en el orden en que aparecen en el texto, los distintos
    elementos/accesorios (sin repetir) que el inspector nombró."""
    vistos_norm = set()
    encontrados = []
    for m in RE_ELEMENTO.finditer(texto):
        valor = m.group(0)
        norm = valor.lower()
        if norm in vistos_norm:
            continue
        vistos_norm.add(norm)
        encontrados.append(valor)
    return encontrados


def _formatear_elementos(lista):
    if not lista:
        return None
    if len(lista) == 1:
        return lista[0]
    return ", ".join(lista[:-1]) + " y " + lista[-1]


def _extraer_variables_de_texto(texto):
    variables = {}
    m = RE_NPS_ETIQUETADO.search(texto) or RE_NPS.search(texto)
    if m:
        variables["nps"] = m.group(1).strip()
    m_de_total = RE_CANTIDAD_DE_TOTAL.search(texto)
    if m_de_total:
        variables["cantidad"] = str(int(m_de_total.group(1)))
    # Puede haber dos cantidades distintas en el mismo comentario (p.ej.
    # "3 uniones bridadas... (04) elementos de ajuste cada una"): la
    # primera es "cantidad" (uniones/uniformes/válvulas...), la segunda,
    # si existe, es "cantidad_elementos" (espárragos/tuercas por unión).
    matches_cantidad = list(RE_CANTIDAD.finditer(texto))
    if matches_cantidad and not m_de_total:
        variables["cantidad"] = str(int(matches_cantidad[0].group(1) or matches_cantidad[0].group(2)))
    if len(matches_cantidad) > 1:
        variables["cantidad_elementos"] = str(int(matches_cantidad[1].group(1) or matches_cantidad[1].group(2)))
    m = RE_LONGITUD.search(texto)
    if m:
        variables["longitud"] = m.group(1)
    m = RE_TIPO_VALVULA.search(texto)
    if m:
        variables["tipo"] = m.group(1).lower()
    else:
        m = RE_TIPO_SOPORTE.search(texto)
        if m:
            variables["tipo"] = m.group(1)
    elementos = _formatear_elementos(_extraer_elementos(texto))
    if elementos:
        variables["elemento"] = elementos
    m = RE_TRAMO_UBICACION.search(texto)
    if m:
        variables["ubicacion"] = f"el tramo {m.group(1).lower()}"
    else:
        m = RE_UBICACION_GENERICA.search(texto)
        if m:
            variables["ubicacion"] = m.group(1).strip()
    return variables


# REGLAS_SUGERENCIA (caso_id, categorías del checklist a las que aplica
# -- substring en minúsculas, o (None,) = cualquiera --, grupos de
# palabras clave -- se exige AL MENOS una coincidencia de CADA grupo --,
# y opcionalmente un 4to elemento de palabras EXCLUIDAS) se carga arriba
# junto con CATALOGO, desde la hoja "Reglas_Sugerencia" del Excel.
# Reglas derivadas 1:1 de COMPLEMENTO/ROL_Y_OBJETIVO.REV2.txt.

# Comentarios que NO representan un hallazgo real (informativos / sin
# incidencia): nunca se sugiere recomendación para ellos.
TEXTOS_SIN_HALLAZGO = (
    "no aplica",
    "sin indicaciones relevantes",
    "sin incidencias",
    "se limita la inspecc",
)

# Regla 1 de ROL_Y_OBJETIVO.REV2.txt: corrosión LEVE y PUNTUAL/LOCALIZADA,
# el componente conserva su integridad -> NUNCA se emite recomendación
# técnica (se clasifica como aceptable). Se detecta por texto, no por caso.
_RE_LEVE_PUNTUAL = re.compile(
    r"\bleve\b.{0,25}\b(puntual|localizad)|(\bno compromete\b|\bconserva su integridad\b|\bsin afectar\b)",
    re.IGNORECASE,
)
TEXTO_ACEPTABLE_SIN_ACCION = (
    "Aceptable — corrosión leve puntual/localizada que no compromete la "
    "integridad del componente; conforme a la matriz de decisión no "
    "corresponde emitir recomendación técnica ni mantenimiento de pintura."
)


def sugerir_caso_desde_texto(categoria_checklist, comentario):
    """Intenta sugerir, con alta confianza y SIN IA, el o los casos del
    CATALOGO que corresponden a un comentario en texto libre del VT-CHECK
    LIST real (por categoría de componente), aplicando las mismas reglas de
    COMPLEMENTO/ROL_Y_OBJETIVO.REV2.txt. Si el comentario describe más de un
    problema y calzan varias reglas, se devuelven combinadas (separadas por
    acciones, como indica la sección de "componentes mixtos" del ROL).
    Devuelve un Resultado(hallazgo, recomendacion) ya redactado y con las
    variables detectables (NPS, cantidad, longitud) extraídas por expresión
    regular, o None si ninguna regla calza con confianza (el hallazgo queda
    entonces para redacción manual) -- salvo que el texto corresponda a
    corrosión leve puntual/localizada, en cuyo caso se marca "Aceptable, sin
    recomendación" tal como exige la matriz de decisión.
    """
    texto_norm = (comentario or "").strip().lower()
    if not texto_norm or any(t in texto_norm for t in TEXTOS_SIN_HALLAZGO):
        return None

    categoria_norm = (categoria_checklist or "").strip().lower()

    resultados = []
    ids_usados = set()
    for regla in REGLAS_SUGERENCIA:
        # El 4to elemento (palabras EXCLUIDAS) es opcional -- reglas que no
        # lo declaran siguen funcionando igual que antes. Sirve para que,
        # p.ej., un caso de severidad "moderada" no se dispare cuando el
        # comentario en realidad dice "moderada a severa" (contiene ambas
        # palabras): la regla exige "moderada" Y que "severa" NO aparezca.
        if len(regla) == 4:
            caso_id, categorias_aplicables, grupos_palabras, palabras_excluidas = regla
        else:
            caso_id, categorias_aplicables, grupos_palabras = regla
            palabras_excluidas = ()

        if categorias_aplicables != (None,) and not any(
            c in categoria_norm for c in categorias_aplicables
        ):
            continue
        if not all(_contiene_alguna(texto_norm, grupo) for grupo in grupos_palabras):
            continue
        if palabras_excluidas and _contiene_alguna(texto_norm, palabras_excluidas):
            continue

        datos = {"caso": caso_id, **_extraer_variables_de_texto(comentario)}
        try:
            resultado = generar_hallazgo_y_recomendacion(datos)
        except ValueError:
            continue
        if resultado is not None and resultado.caso_id not in ids_usados:
            resultados.append(resultado)
            ids_usados.add(resultado.caso_id)

    if resultados:
        if len(resultados) == 1:
            return resultados[0]
        # Más de un problema detectado en la MISMA observación (p.ej. una
        # válvula con corrosión Y volante roto a la vez): se unifica en UNA
        # sola oración -- la primera recomendación tal cual, y cada una
        # adicional conectada con "Así como"/"Asimismo,"/"Adicionalmente,"
        # (el primero que no repita una frase ya usada en el texto previo)
        # en vez de repetirse como una nueva oración imperativa aparte (sin
        # encabezado "Recomendación:" ni viñetas, que la hacían ver más
        # larga de lo necesario).
        partes = [resultados[0].recomendacion.strip()]
        for r in resultados[1:]:
            conector = _elegir_conector(" ".join(partes))
            partes.append(_como_continuacion(r.recomendacion, conector))
        return Resultado(
            caso_id="+".join(r.caso_id for r in resultados),
            hallazgo=" ".join(r.hallazgo for r in resultados),
            recomendacion=" ".join(partes),
        )

    if _RE_LEVE_PUNTUAL.search(texto_norm):
        return Resultado(
            caso_id="ACEPTABLE_SIN_RECOMENDACION",
            hallazgo=comentario.strip(),
            recomendacion=TEXTO_ACEPTABLE_SIN_ACCION,
        )

    return None


def aplicar_caso_manual(caso_id, comentario):
    """Aplica un caso del CATALOGO elegido explícitamente (nunca por
    palabras clave) a un comentario de campo -- pensado para cuando
    sugerir_caso_desde_texto() no encuentra ninguna regla con confianza
    (hallazgo marcado PENDIENTE) y una tercera parte (persona o un agente de
    IA leyendo el propio COMPENDIO) decide cuál de los casos YA APROBADOS
    del catálogo corresponde. El hallazgo y la recomendación finales siguen
    saliendo EXACTOS del catalogo_tpl -- igual que con la sugerencia
    automática, nunca se redacta texto nuevo ni se inventa una
    recomendación: solo cambia QUIÉN elige el caso, nunca de dónde sale el
    texto. Lanza ValueError si caso_id no existe en el catálogo (mismo
    comportamiento que generar_hallazgo_y_recomendacion)."""
    datos = {"caso": caso_id, **_extraer_variables_de_texto(comentario)}
    return generar_hallazgo_y_recomendacion(datos)


# ==============================================================================
# MECANISMOS DE DAÑO (catálogo fijo de 20, cruzado por palabras clave contra
# los hallazgos detectados en la inspección -- sin IA).
# ==============================================================================
CATALOGO_MECANISMOS_DANO = [
    "Corrosión Atmosférica",
    "Corrosión Galvánica",
    "Corrosión por Suelo",
    "Corrosión de la célula de concentración",
    "Punto de contacto",
    "Fatiga Mecánica",
    "Erosión/Corrosión Erosión",
    "Corrosión por H2S",
    "Corrosión Inducida por Microorganismos (Mic)",
    "Corrosión por CO2",
    "Corrosión Bajo Aislamiento (CUI)",
    "Corrosión Bajo Ignifugado (CUF)",
    "Corrosión por Agua de Condensado de Caldera",
    "Corrosión Caustica",
    "Sulfuración",
    "SCC por Cloruro",
    "SCC por Cáusticos",
    "SCC por Amina",
    "Corrosión Por Ácido Clorhídrico (HCl)",
    "Corrosión Por Agua Amarga (Ácida)",
]

# Cada mecanismo: patrones regex (con límites de palabra donde hace falta,
# para evitar falsos positivos con siglas cortas como "mic"/"cui"/"cuf"/"hcl").
_PATRONES_MECANISMO = {
    "Corrosión Atmosférica": (r"corrosi[oó]n atmosf[eé]rica",),
    "Corrosión Galvánica": (r"galv[aá]nic",),
    "Corrosión por Suelo": (r"corrosi[oó]n por suelo", r"tuber[ií]a enterrada", r"soterrad"),
    "Corrosión de la célula de concentración": (r"c[eé]lula de concentraci[oó]n",),
    "Punto de contacto": (r"punto de contacto", r"contacto met[aá]l-?met[aá]l", r"contacto directo"),
    "Fatiga Mecánica": (r"fatiga mec[aá]nica",),
    "Erosión/Corrosión Erosión": (r"erosi[oó]n",),
    "Corrosión por H2S": (r"\bh2s\b", r"[aá]cido sulfh[ií]drico"),
    "Corrosión Inducida por Microorganismos (Mic)": (r"microorganismos", r"\bmic\b"),
    "Corrosión por CO2": (r"\bco2\b", r"di[oó]xido de carbono"),
    "Corrosión Bajo Aislamiento (CUI)": (r"bajo aislamiento", r"\bcui\b", r"aislamiento t[eé]rmico"),
    "Corrosión Bajo Ignifugado (CUF)": (r"ignif[uú]g", r"\bcuf\b"),
    "Corrosión por Agua de Condensado de Caldera": (r"agua de condensado", r"condensado de caldera"),
    "Corrosión Caustica": (r"c[aá]ustic",),
    "Sulfuración": (r"sulfuraci[oó]n",),
    "SCC por Cloruro": (r"scc.{0,15}cloruro", r"cloruro.{0,15}scc"),
    "SCC por Cáusticos": (r"scc.{0,15}c[aá]ustic", r"c[aá]ustic.{0,15}scc"),
    "SCC por Amina": (r"scc.{0,15}amina", r"amina.{0,15}scc"),
    "Corrosión Por Ácido Clorhídrico (HCl)": (r"\bhcl\b", r"[aá]cido clorh[ií]drico"),
    "Corrosión Por Agua Amarga (Ácida)": (r"agua amarga",),
}
_PATRONES_MECANISMO_COMPILADOS = {
    m: [re.compile(p, re.IGNORECASE) for p in pats] for m, pats in _PATRONES_MECANISMO.items()
}

# Reglas de respaldo confirmadas con el usuario: un hallazgo que menciona
# "corrosión" en términos genéricos, sin calificarla con ninguna palabra
# clave de las de arriba, se asume por defecto Corrosión Atmosférica.
_RE_CORROSION_GENERICA = re.compile(r"corrosi[oó]n", re.IGNORECASE)

# Asociación por el PRODUCTO/FLUIDO de la línea (columna "NOMBRE DEL FLUIDO"
# de la base maestra), además de por el texto del hallazgo -- confirmado con
# el usuario. Deliberadamente conservador: solo se mapean servicios con una
# correlación de mecanismo bien establecida (API 571); los hidrocarburos
# genéricos (gasolina, diésel, GLP, etc., sin más calificación) NO se
# asocian a ningún mecanismo por el solo nombre del fluido.
_PATRONES_MECANISMO_FLUIDO = {
    "Corrosión por H2S": (r"h2s", r"[aá]cido sulfh[ií]drico", r"\bsour\b", r"\bagria\b"),
    "SCC por Amina": (r"amina",),
    "Corrosión Por Agua Amarga (Ácida)": (r"agua\s*[aá]cida", r"aguas?\s*[aá]cidas"),
    "Corrosión Caustica": (r"c[aá]ustic", r"\bsosa\b"),
    "SCC por Cloruro": (r"clorur",),
    "Corrosión Por Ácido Clorhídrico (HCl)": (r"\bhcl\b", r"[aá]cido clorh[ií]drico"),
    "Corrosión por CO2": (r"\bco2\b", r"di[oó]xido de carbono"),
    "Corrosión por Agua de Condensado de Caldera": (r"condensado", r"agua.{0,15}caldera"),
}
_PATRONES_MECANISMO_FLUIDO_COMPILADOS = {
    m: [re.compile(p, re.IGNORECASE) for p in pats]
    for m, pats in _PATRONES_MECANISMO_FLUIDO.items()
}


def detectar_mecanismos_dano(textos_hallazgos, fluidos=None):
    """Cruza los hallazgos (texto libre, mejorado) del grupo -- y, si se
    provee, el nombre del fluido/producto de cada línea (columna "NOMBRE DEL
    FLUIDO" de la base maestra) -- contra el catálogo fijo de 20 mecanismos
    de daño por palabras clave (sin IA), y devuelve la lista de mecanismos
    que efectivamente aplican, en el orden del catálogo.

    Reglas de respaldo (confirmadas con el usuario):
    - Un hallazgo que menciona "corrosión" sin calificarla con ninguna
      palabra clave específica se asume Corrosión Atmosférica por defecto.
    - El fluido/producto de la línea también puede asociar un mecanismo
      (p.ej. servicio con amina -> SCC por Amina), independientemente de lo
      que diga el texto del hallazgo."""
    texto_total = " ".join(t or "" for t in textos_hallazgos)
    detectados = set()
    for mecanismo in CATALOGO_MECANISMOS_DANO:
        patrones = _PATRONES_MECANISMO_COMPILADOS.get(mecanismo, ())
        if any(p.search(texto_total) for p in patrones):
            detectados.add(mecanismo)

    for texto in textos_hallazgos:
        texto = texto or ""
        if not _RE_CORROSION_GENERICA.search(texto):
            continue
        calificada = any(
            p.search(texto)
            for patrones in _PATRONES_MECANISMO_COMPILADOS.values()
            for p in patrones
        )
        if not calificada:
            detectados.add("Corrosión Atmosférica")
            break

    for fluido in (fluidos or []):
        fluido = fluido or ""
        for mecanismo, patrones in _PATRONES_MECANISMO_FLUIDO_COMPILADOS.items():
            if any(p.search(fluido) for p in patrones):
                detectados.add(mecanismo)

    return [m for m in CATALOGO_MECANISMOS_DANO if m in detectados]
