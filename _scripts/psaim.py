"""
psaim.py — lectura del reporte PSAIM (Siemens Energy) y cálculo de Rate de
Corrosión y Vida Útil, según las reglas confirmadas en las secciones 9.2,
10.2 y 13.3 del análisis:

    Rate de corrosión (mm/año) = RCR (MPY, dato de cabecera) × 0.0254,
        redondeado a 3 decimales (NO a precisión completa, aunque el PSAIM
        traiga más — confirmado explícitamente en 10.2 pese a que el
        informe de referencia GT-032 mostraba precisión completa).

    Vida remanente (años) = MÍNIMO de la columna "TML Vida útil" (detalle
        por punto de medición) — NUNCA el "Rest. Vida útil (desde el
        último relevamiento...)" del resumen de cabecera, que puede ser
        más optimista que el peor punto real.

⚠️ Solo sirve la versión del PSAIM que trae la tabla de detalle por TML
(columna "TML Vida útil" fila por fila) — la versión resumen tipo
"... DATOS.xlsx" no trae esa columna y no se puede usar para este cálculo
(sección 9.2 / checklist de insumos, sección 2).
"""
import re
import unicodedata
import openpyxl

MPY_TO_MMYEAR = 0.0254

# Umbral de vida útil por Clase API 570 para decidir si se muestra el valor
# exacto o ">umbral" en el SUMARIO (sección 9.7a, ratificado en 10.2).
UMBRAL_POR_CLASE = {
    "clase 1": 10,
    # todas las demás clases (2, 3, 4, punto de inyección, etc.) -> 20
}
UMBRAL_DEFAULT = 20


def _norm(s):
    if s is None:
        return ""
    s = re.sub(r"[\s\xa0]+", "", str(s)).lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s


class PSAIMFaltaDetalle(Exception):
    """La versión del archivo no trae la tabla de detalle por TML."""


def leer_psaim(path):
    """Devuelve {'rcr_mpy': float, 'vida_util_anios': float (mínimo de TML
    Vida útil), 'n_tml': int}. Lanza PSAIMFaltaDetalle si no encuentra la
    columna de detalle."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rcr = None
    header_row_idx = None
    vida_col_idx = None

    all_rows = list(ws.iter_rows(values_only=True))

    # 1) Buscar RCR = <num> MPY en cualquier celda (texto libre, cabecera).
    for row in all_rows:
        for cell in row:
            if isinstance(cell, str) and "RCR" in cell.upper():
                m = re.search(r"RCR\s*=\s*([\d.,]+)\s*MPY", cell, re.IGNORECASE)
                if m:
                    rcr = float(m.group(1).replace(",", "."))
                    break
        if rcr is not None:
            break

    # 2) Buscar la fila de encabezado de la tabla de detalle (contiene una
    # celda cuyo texto normalizado sea "tmlvidautil").
    for r_idx, row in enumerate(all_rows):
        for c_idx, cell in enumerate(row):
            if "tmlvidautil" in _norm(cell):
                header_row_idx = r_idx
                vida_col_idx = c_idx
                break
        if header_row_idx is not None:
            break

    if rcr is None:
        raise PSAIMFaltaDetalle(f"No se encontró 'RCR = ... MPY' en {path}")
    if header_row_idx is None:
        raise PSAIMFaltaDetalle(
            f"No se encontró la columna 'TML Vida útil' en {path} — "
            "probablemente es la versión resumen ('... DATOS.xlsx'), "
            "que no trae la tabla de detalle por TML. Se necesita la "
            "versión con detalle (ej. '..._rev_FECHA.xlsx')."
        )

    valores = []
    for row in all_rows[header_row_idx + 1:]:
        if vida_col_idx >= len(row):
            continue
        v = row[vida_col_idx]
        if isinstance(v, (int, float)):
            valores.append(float(v))

    if not valores:
        raise PSAIMFaltaDetalle(f"La columna 'TML Vida útil' está vacía en {path}")

    return {
        "rcr_mpy": rcr,
        "vida_util_anios": min(valores),
        "n_tml": len(valores),
    }


def rate_corrosion_mm_anio(rcr_mpy):
    return round(rcr_mpy * MPY_TO_MMYEAR, 3)


def umbral_clase(clase_api570):
    return UMBRAL_POR_CLASE.get(str(clase_api570).strip().lower(), UMBRAL_DEFAULT)


def vida_util_display(vida_util_anios, clase_api570):
    """Sección 9.7a: si la vida útil calculada es MENOR al umbral de su
    Clase, se muestra el valor exacto (ej. "8.3"); si es igual o mayor, se
    muestra ">umbral" (ej. ">20"), nunca el valor preciso."""
    umbral = umbral_clase(clase_api570)
    if vida_util_anios < umbral:
        return f"{vida_util_anios:g}"
    return f">{umbral}"
