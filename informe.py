"""
informe.py — ensambla el .docx final de un grupo a partir de la plantilla
recuperada (`assets/plantilla_base.docx`, en realidad el último informe real
entregado — GT-010 — que ya trae todas las correcciones de formato de las
secciones 11.5/11.6 aplicadas) y los datos ya cruzados/calculados del grupo
nuevo (inventario.py + psaim.py + checklist.py + el JSON de
hallazgos/recomendaciones que Claude redacta cada sesión con el COMPENDIO).

Técnica general: nunca se reconstruye el documento desde cero — siempre se
PATCHEA la plantilla real (regla de la sección 6/7, aprendida tras el
rechazo del usuario a un informe armado con docx-js). Ver docxlib.py para
los helpers de bajo nivel.

Anclas usadas para ubicar los párrafos de un solo ejemplar (no clonados por
línea) dentro de la plantilla — verificadas contra el .docx real de GT-010
el 15/09/2026, con fallback por contenido de texto si algún día cambia la
plantilla base:
    reco1            w14:paraId 5A106A4B  /  empieza con "Programar la próxima"
    nota VT          w14:paraId 1B20563E  /  empieza con "En la"
    fecha/personal   w14:paraId 0BDF29A0  /  empieza con "La Inspección se realizó"
"""
import re
import docx
import docxlib
from docxlib import qn

# --- Constantes de negocio (secciones 9.3, 9.4, 9.7b) -----------------------

API570_INTERVALOS = {
    "clase 1": (
        "Programar la próxima inspección visual y medición de espesores {obj} "
        "en un periodo no mayor a 5 años."
    ),
    "clase 2": (
        "Programar la próxima inspección visual en un periodo no mayor a 5 "
        "años y medición de espesores {obj} en un intervalo máximo de 10 "
        "años."
    ),
    "clase 3": (
        "Programar la próxima inspección visual y medición de espesores {obj} "
        "en un periodo no mayor a 10 años."
    ),
    "clase 4": (
        "Programar la próxima medición de espesores {obj} en un periodo no "
        "mayor a 10 años y la próxima inspección externa en un periodo no "
        "mayor a 5 años. (opcional por ser tubería clase 04)."
    ),
    "punto de inyeccion": (
        "Programar la próxima inspección visual y medición de espesores {obj} "
        "en un periodo no mayor a 3 años."
    ),
}
CLASE_LABEL = {
    "clase 1": "Clase 1", "clase 2": "Clase 2",
    "clase 3": "Clase 3", "clase 4": "Clase 4",
    "punto de inyeccion": "Punto de inyección",
}
NOTA_PUNTO_INYECCION = (
    "Monitorear y ampliar la zona de examinación de los puntos de inyección "
    "en la frecuencia determinada (3 años), considerando las áreas de "
    "examinación corriente arriba y corriente abajo recomendadas por el "
    "estándar API 570 párrafo 5.10."
)

MECANISMOS_DANO_VALIDOS = [
    "Corrosión Atmosférica", "Corrosión Galvánica", "Corrosión por Suelo",
    "Corrosión de la célula de concentración", "Punto de contacto",
    "Fatiga Mecánica", "Erosión/Corrosión", "Corrosión por H2S",
    "Corrosión Inducida por Microorganismos (Mic)", "Corrosión por CO2",
    "Corrosión Bajo Aislamiento (CUI)", "Corrosión Bajo Ignifugado (CUF)",
    "Corrosión por Agua de Condensado de Caldera", "Corrosión Cáustica",
    "Sulfuración", "SCC por Cloruro", "SCC por Cáusticos", "SCC por Amina",
    "Corrosión por Ácido Clorhídrico (HCl)", "Corrosión por Agua Amarga (Ácida)",
]

PARA_ID_RECO1 = "5A106A4B"
PARA_ID_NOTA_VT = "1B20563E"
PARA_ID_FECHA_PERSONAL = "0BDF29A0"

SIN_DATO = "SIN DATO"


def _clase_key(clase_txt):
    return re.sub(r"\s+", " ", str(clase_txt or "")).strip().lower()


def build_reco1_textos(grupo, lineas_por_clase):
    """`lineas_por_clase`: dict {clase_normalizada: [N° de ítem...]}.

    REGLA CONFIRMADA (18/09/2026, validada por el usuario con 2 casos
    reales -- 22-GLP-GT-023 mixto y 26-FLARE-GT-029 uniforme):

    - Si TODAS las líneas del grupo comparten la misma Clase API 570: UN
      SOLO párrafo, sin mencionar Clase ni numeración, con "del grupo de
      tuberías <GRUPO>". Sin coletilla de tabla/CML's. Ejemplo real
      (26-FLARE-GT-029, 100% Clase 1):
          "Programar la próxima inspección visual y medición de espesores
           del grupo de tuberías 26-FLARE-GT-029 en un periodo no mayor a
           5 años."

    - Si el grupo MEZCLA 2 o más Clases: UN párrafo POR Clase, cada uno
      mencionando la numeración de ítem (nunca el TAG completo -- misma
      convención de la nota "solo VT"), la Clase entre paréntesis, Y el
      grupo de tuberías. Ejemplo real (22-GLP-GT-023, Clase 1 + Clase 2):
          "Programar la próxima inspección visual y medición de espesores
           de las líneas N° 1, 2, 4, 7, 8, 9 y 10 (Clase 1) del grupo de
           tuberías 22-GLP-GT-023 en un periodo no mayor a 5 años."
          "Programar la próxima inspección visual en un periodo no mayor
           a 5 años y medición de espesores de las líneas N° 3, 5, 6, 11,
           12 y 13 (Clase 2) del grupo de tuberías 22-GLP-GT-023 en un
           intervalo máximo de 10 años."

    Estas 2 ramas (uniforme vs. mixto) son las únicas variantes -- no
    inventar un tercer formato intermedio.
    """
    clases = [c for c in lineas_por_clase if c in API570_INTERVALOS]
    textos = []
    if len(clases) == 1:
        obj = f"del grupo de tuberías {grupo}"
        textos.append(API570_INTERVALOS[clases[0]].format(obj=obj))
    else:
        for clase in clases:
            nums = lineas_por_clase[clase]
            label = CLASE_LABEL.get(clase, clase.title())
            if len(nums) == 1:
                obj = (f"de la línea N° {nums[0]} ({label}) del grupo de "
                       f"tuberías {grupo}")
            else:
                nums_str = ", ".join(str(n) for n in nums[:-1]) + f" y {nums[-1]}"
                obj = (f"de las líneas N° {nums_str} ({label}) del grupo de "
                       f"tuberías {grupo}")
            textos.append(API570_INTERVALOS[clase].format(obj=obj))
    if "punto de inyeccion" in lineas_por_clase:
        textos.append(NOTA_PUNTO_INYECCION)
    return textos


def build_nota_vt_texto(lineas):
    """Sección 11.5.2: nota bajo la tabla SUMARIO listando (por NÚMERO de
    ítem, no de TAG) las líneas con alcance VT-CIRCUITOS ("solo visual").
    Si el grupo no mezcla alcances (todas o ninguna con UT), el párrafo no
    aplica (se devuelve None para que el llamador lo borre)."""
    solo_vt = [i + 1 for i, ln in enumerate(lineas)
               if str(ln.get("alcance", "")).upper() != "LINEAS"]
    if not solo_vt or len(solo_vt) == len(lineas):
        return None
    if len(solo_vt) == 1:
        ref = f"la línea N° {solo_vt[0]}"
        verbo = "se realizó"
    else:
        nums = ", ".join(str(n) for n in solo_vt[:-1]) + f" y {solo_vt[-1]}"
        ref = f"las líneas N° {nums}"
        verbo = "se realizó"
    return (f"En {ref}, solo {verbo} inspección visual, tal como se "
            "especifica en los alcances del servicio. ")


def build_circuito_linea_texto(linea, cruzo_inventario, psaim_data, clase,
                                hallazgo_texto=None):
    """Réplica de las 5 variantes confirmadas en el .docx real de GT-010
    (sección de validación, 15/09/2026) para la columna 'HALLAZGOS
    RELEVANTES EN VT Y UT' de la tabla 9.0, según qué insumo falte."""
    sin_fecha = not linea.get("fecha_inspeccion")
    tag_ambiguo = not cruzo_inventario

    if sin_fecha:
        return "PENDIENTE (línea aún sin inspección de campo)"
    # Nota: tag_ambiguo (TAG sin coincidencia en la base de datos técnica
    # maestra, típico de líneas "NN-xx" aún sin codificación oficial) ya NO
    # pisa el hallazgo real de campo -- solo afecta la tabla 7.0 (cruce
    # técnico), no la 9.0 ni la 2.0, porque el checklist VT sí existe e
    # inspeccionó la línea aunque le falte TAG oficial en el inventario.

    prefijo = None
    if str(linea.get("alcance", "")).upper() == "LINEAS":
        if psaim_data:
            rate = psaim_data["rate"]
            vida = psaim_data["vida_display"]
            prefijo = f"Rate de corrosión {rate} mm/año, con una vida remanente {vida} años."
        else:
            prefijo = ("Rate de corrosión y vida remanente: PENDIENTE "
                       "(falta el reporte PSAIM de esta línea).")

    if isinstance(hallazgo_texto, (list, tuple)):
        items = list(hallazgo_texto)
    elif hallazgo_texto:
        items = [hallazgo_texto]
    else:
        items = ["PENDIENTE (falta checklist VT)"]

    if prefijo is not None:
        return ("__PREFIJO_MAS_LISTA__", prefijo, items)
    return items if len(items) > 1 else items[0]


def build_recomendacion_texto(recomendacion_texto=None):
    if isinstance(recomendacion_texto, (list, tuple)):
        items = list(recomendacion_texto)
        return items if len(items) > 1 else (items[0] if items else
               "PENDIENTE (falta checklist VT para redactar hallazgo/recomendación)")
    return recomendacion_texto or (
        "PENDIENTE (falta checklist VT para redactar hallazgo/recomendación)"
    )


# --- Localización de párrafos ancla ------------------------------------------

def _find_by_id_or_prefix(document, para_id, prefix):
    p = docxlib.find_paragraph_by_id(document, para_id)
    if p is not None:
        return p
    for cand in document.element.body.iter(qn("w:p")):
        txt = "".join(t.text or "" for t in cand.findall(f'.//{qn("w:t")}'))
        if txt.strip().startswith(prefix):
            return cand
    return None


def _replace_grupo_paragraphs(document, grupo_nuevo):
    """Reemplaza el título 'GRUPO <código>' de portada — SOLO párrafos de
    cuerpo (document.paragraphs excluye, a propósito, los que están dentro
    de tablas: la celda de la tabla portada 'GRUPO DE TUBERÍAS
    INSPECCIONADAS' también empieza con 'GRUPO ' y NO debe tocarse aquí —
    esa la maneja _patch_portada_tabla por separado)."""
    for p in document.paragraphs:
        txt = p.text.strip()
        if txt.startswith("GRUPO ") and not txt.upper().startswith("GRUPO DE"):
            docxlib.set_para_content(p._p, f"GRUPO {grupo_nuevo}")


def _replace_banner_codigo(document, codigo_informe):
    """codigo_informe: solo la parte variable, ej. '1018-2026' (sin el
    prefijo fijo 'ADEMINSAC-FIAB-RLP-')."""
    nuevo_texto = f"ADEMINSAC-FIAB-RLP-{codigo_informe}"
    encontrados = 0
    for p in document.element.body.iter(qn("w:p")):
        direct_runs = p.findall(qn("w:r"))
        txt = "".join(t.text or "" for r in direct_runs for t in r.findall(qn("w:t")))
        if txt.startswith("ADEMINSAC-FIAB-RLP-"):
            docxlib.set_para_content(p, nuevo_texto, remove_highlight=False)
            encontrados += 1
    return encontrados


def _patch_portada_tabla(document, grupo, fecha_ini, fecha_fin):
    """Tabla 2 (portada): celda 'GRUPO DE TUBERÍAS INSPECCIONADAS' y celda
    de 'FECHA DE INSPECCIÓN' (rango, sección 13.7)."""
    for table in document.tables:
        for row in table.rows:
            cells = row.cells
            for i, cell in enumerate(cells):
                t = cell.text.strip().upper()
                if "GRUPO DE TUBERÍAS INSPECCIONADAS" in t and i + 1 < len(cells):
                    docxlib.set_cell_text(cells[i + 1]._tc, grupo)
        # fila con 'FECHA DE INSPECCIÓN' encabezado y el rango en la fila
        # siguiente (mismo layout que el real: header row + value row)
        for r_idx in range(len(table.rows) - 1):
            row_cells = table.rows[r_idx].cells
            for i, cell in enumerate(row_cells):
                if "FECHA DE INSPECCIÓN" in cell.text.upper():
                    val_cell = table.rows[r_idx + 1].cells[i]
                    docxlib.set_cell_text(val_cell._tc, f"{fecha_ini} al {fecha_fin}")


def _patch_reco1(document, grupo, lineas):
    p = _find_by_id_or_prefix(document, PARA_ID_RECO1, "Programar la próxima")
    if p is None:
        return ["No se encontró el párrafo de Recomendación #1 (reco1) en la plantilla."]
    lineas_por_clase = {}
    for i, ln in enumerate(lineas, start=1):
        clase = _clase_key(ln.get("clase"))
        lineas_por_clase.setdefault(clase, []).append(i)
    textos = build_reco1_textos(grupo, lineas_por_clase)
    if not textos:
        return ["No se pudo determinar el texto de Recomendación #1: "
                "ninguna línea tiene una Clase API 570 reconocida."]
    docxlib.clone_paragraph_block(p, textos)
    return []


def _patch_nota_vt(document, lineas):
    p = _find_by_id_or_prefix(document, PARA_ID_NOTA_VT, "En la")
    if p is None:
        return ["No se encontró el párrafo de nota 'solo VT' en la plantilla."]
    texto = build_nota_vt_texto(lineas)
    if texto is None:
        parent = p.getparent()
        parent.remove(p)
    else:
        docxlib.set_para_content(p, texto)
    return []


def _patch_fecha_personal(document, fecha_ini, fecha_fin, examinadores):
    p = _find_by_id_or_prefix(document, PARA_ID_FECHA_PERSONAL, "La Inspección se realizó")
    if p is None:
        return ["No se encontró el párrafo de fecha/personal en la plantilla."]
    docxlib.set_para_content(
        p,
        f"La Inspección se realizó {fecha_ini} al {fecha_fin}. El personal "
        "técnico de ADEMINSAC que participó, fue el siguiente: ",
    )

    # Bloque de examinadores: los 3 primeros (Jefe de Proyecto / Inspector
    # API 570 / Supervisor de Equipos Estáticos) son SIEMPRE fijos y nunca
    # se tocan (sección 12.2); el resto (a partir del 4to párrafo tras la
    # fecha) es el bloque variable a reemplazar por `examinadores`.
    siguiente = p
    fijos = 0
    variables_anchor = None
    while fijos < 3:
        siguiente = siguiente.getnext()
        if siguiente is None:
            return ["No se encontraron los 3 examinadores fijos tras el "
                    "párrafo de fecha/personal."]
        fijos += 1
    variables_anchor = siguiente.getnext()
    if variables_anchor is None:
        return ["No se encontró el bloque de examinadores variable."]

    # Recolectar y borrar todos los párrafos variables existentes (los que
    # matchean el patrón "Nombre: (Examinador Nivel II...")
    to_remove = []
    cursor = variables_anchor
    while cursor is not None:
        txt = "".join(t.text or "" for t in cursor.findall(f'.//{qn("w:t")}'))
        if re.search(r"\(Examinador Nivel", txt):
            to_remove.append(cursor)
            cursor = cursor.getnext()
        else:
            break

    if not to_remove:
        return ["No se encontró el bloque variable de examinadores para reemplazar."]

    anchor = to_remove[0]
    for extra in to_remove[1:]:
        extra.getparent().remove(extra)
    docxlib.clone_paragraph_block(anchor, examinadores)
    return []


def _patch_mecanismo_dano(document, mecanismos):
    table = None
    for t in document.tables:
        if t.rows[0].cells[0].text.strip().upper() == "ITEM" and \
           "MECANISMO DE DAÑO" in t.rows[0].cells[1].text.upper():
            table = t
            break
    if table is None:
        return ["No se encontró la tabla de Mecanismo de Daño (8.0)."]
    if not mecanismos:
        return []
    rows = docxlib.clone_table_to_n_rows(table, len(mecanismos), header_rows=1)
    for i, (row, mec) in enumerate(zip(rows, mecanismos), start=1):
        docxlib.fill_row(row, [str(i), mec, ""])
    return []


def _patch_portada_foto(out_path, foto_path):
    info = docxlib.find_portada_photo(out_path)
    if info is None:
        return ["No se pudo identificar la imagen de portada para reemplazarla."]
    media_name, w, h = info
    media_name = media_name.split("/")[-1]  # find_portada_photo ya devuelve
    # la ruta completa "word/media/xxx"; replace_media_in_docx antepone el
    # prefijo "word/media/" de nuevo -- pasar solo el nombre de archivo.
    with open(foto_path, "rb") as f:
        data = f.read()
    docxlib.replace_media_in_docx(out_path, out_path, media_name, data)
    return []


# --- Tablas por línea ---------------------------------------------------

def _find_tabla_by_headers(document, header_texts, n_cols=None):
    for t in document.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if n_cols and len(hdr) != n_cols:
            continue
        if all(any(h.upper().startswith(exp.upper()) for h in hdr) for exp in header_texts):
            return t
    return None


def patch_informe(config, plantilla_path, out_path):
    """Función principal: arma el .docx final del grupo.

    `config` (ya resuelto por generar_informe.py):
        grupo, unidad, codigo_informe (ej. '1018-2026'),
        fecha_ini, fecha_fin, examinadores (lista de strings ya formateados
        "Nombre: (Examinador Nivel II...)"),
        lineas: lista de dicts con al menos:
            item, tag, sap, alcance, fecha_inspeccion,
            clase, cruce (dict de inventario.cruzar_linea o None),
            psaim (dict {rate, vida_display} o None),
            hallazgo (str o None), recomendacion (str o None)
        mecanismos_dano: lista de strings (subset de MECANISMOS_DANO_VALIDOS)
        foto_portada: path opcional

    Devuelve lista de avisos (strings) — nunca lanza excepción por un
    insumo faltante, salvo error estructural grave (regla 11.7: mejor
    avisar y seguir que detener todo el pipeline).
    """
    avisos = []
    document = docx.Document(plantilla_path)

    grupo = config["grupo"]
    lineas = config["lineas"]
    n = len(lineas)

    _replace_grupo_paragraphs(document, grupo)
    encontrados = _replace_banner_codigo(document, config["codigo_informe"])
    if encontrados == 0:
        avisos.append("No se encontró el banner de portada con el código de informe.")

    _patch_portada_tabla(document, grupo, config["fecha_ini"], config["fecha_fin"])
    avisos += _patch_reco1(document, grupo, lineas)
    avisos += _patch_nota_vt(document, lineas)
    avisos += _patch_fecha_personal(document, config["fecha_ini"], config["fecha_fin"],
                                     config["examinadores"])
    avisos += _patch_mecanismo_dano(document, config.get("mecanismos_dano", []))

    # SUMARIO (1.0): N°, SAP, Línea, Rate Corrosión, Vida Útil
    t_sumario = _find_tabla_by_headers(document, ["N°", "SAP"], n_cols=5)
    if t_sumario is None:
        avisos.append("No se encontró la tabla SUMARIO (1.0).")
    else:
        rows = docxlib.clone_table_to_n_rows(t_sumario, n, header_rows=1)
        for i, (row, ln) in enumerate(zip(rows, lineas), start=1):
            if str(ln.get("alcance", "")).upper() == "LINEAS" and ln.get("psaim"):
                rate, vida = ln["psaim"]["rate"], ln["psaim"]["vida_display"]
            else:
                rate, vida = "---", "---"
            docxlib.fill_row(row, [str(i), ln.get("sap", SIN_DATO), ln["tag"], rate, vida])

    # RECOMENDACIONES (2.0): Ítem, Línea, Recomendación
    t_reco = _find_tabla_by_headers(document, ["Ítem", "Línea", "Recomendación"], n_cols=3)
    if t_reco is None:
        avisos.append("No se encontró la tabla RECOMENDACIONES (2.0).")
    else:
        rows = docxlib.clone_table_to_n_rows(t_reco, n, header_rows=1)
        for i, (row, ln) in enumerate(zip(rows, lineas), start=1):
            texto = build_recomendacion_texto(ln.get("recomendacion"))
            docxlib.fill_row_multi(row, [str(i), ln["tag"], texto])

    # 7.0 Características de la línea (13 columnas, 2 filas de encabezado)
    t_70 = _find_tabla_by_headers(document, ["ITEM", "SAP", "TAG"], n_cols=13)
    if t_70 is None:
        avisos.append("No se encontró la tabla 7.0 (Características de la línea).")
    else:
        rows = docxlib.clone_table_to_n_rows(t_70, n, header_rows=2)
        for i, (row, ln) in enumerate(zip(rows, lineas), start=1):
            c = ln.get("cruce") or {}
            vals = [
                str(i), ln.get("sap", SIN_DATO), ln["tag"],
                c.get("pres_oper_psi", SIN_DATO), c.get("temp_oper_f", SIN_DATO),
                c.get("pres_dis_psi", SIN_DATO), c.get("temp_dis_f", SIN_DATO),
                c.get("schedule", SIN_DATO), c.get("material", SIN_DATO),
                c.get("inicio", SIN_DATO), c.get("termino", SIN_DATO),
                c.get("fluido", SIN_DATO), c.get("clase", SIN_DATO),
            ]
            docxlib.fill_row(row, vals)

    # 9.0 CIRCUITO/LÍNEA INTERVENIDA: ITEM, UNIDAD, LÍNEA, HALLAZGOS
    t_circ = _find_tabla_by_headers(document, ["ITEM", "UNIDAD", "LÍNEA"], n_cols=4)
    if t_circ is None:
        avisos.append("No se encontró la tabla 9.0 (CIRCUITO/LÍNEA INTERVENIDA).")
    else:
        rows = docxlib.clone_table_to_n_rows(t_circ, n, header_rows=1)
        for i, (row, ln) in enumerate(zip(rows, lineas), start=1):
            texto = build_circuito_linea_texto(
                ln, ln.get("cruce") is not None, ln.get("psaim"),
                ln.get("clase"), ln.get("hallazgo"),
            )
            tcs = row.findall(docxlib.qn("w:tc"))
            docxlib.set_cell_text(tcs[0], str(i))
            docxlib.set_cell_text(tcs[1], str(config.get("unidad", SIN_DATO)))
            docxlib.set_cell_text(tcs[2], ln["tag"])
            if isinstance(texto, tuple) and texto and texto[0] == "__PREFIJO_MAS_LISTA__":
                _, prefijo, items = texto
                docxlib.set_cell_text_prefijo_mas_lista(tcs[3], prefijo, items)
            elif isinstance(texto, (list, tuple)):
                docxlib.set_cell_text_multi(tcs[3], texto)
            else:
                docxlib.set_cell_text(tcs[3], str(texto))

    document.save(out_path)

    if config.get("foto_portada"):
        avisos += _patch_portada_foto(out_path, config["foto_portada"])

    return avisos
