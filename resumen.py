"""Resumen y balance total de los dos planes (sin Streamlit)."""

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (clave en kpis, etiqueta). Todas cuentan informes salvo las de líneas.
FILAS = [
    ("informes", "Informes totales"),
    ("finalizados", "Informes finalizados"),
    ("pend_elaborar", "Pendientes de elaborar"),
    ("pend_inspeccion", "Campo · pend. inspección"),
    ("por_elaborar", "Gabinete · por elaborar"),
    ("en_proceso", "Gabinete · en elaboración"),
    ("por_revisar", "Especialista · por revisar"),
    ("revisados", "Especialista · revisados"),
    ("en_revision", "Cliente · en revisión"),
    ("valorizados", "Cliente · valorizados"),
    ("lineas", "Líneas activas"),
    ("lineas_retiradas", "Líneas retiradas"),
]
# Pendientes que requieren acción, por responsable de la etapa.
ALERTAS = [
    ("pend_inspeccion", "Pend. inspección (campo)"),
    ("por_elaborar", "Por elaborar (gabinete)"),
    ("por_revisar", "Por revisar (especialista)"),
    ("en_revision", "En revisión (cliente)"),
]


def kpis_adicionales(kp):
    """Adapta los KPIs de planes.py (TOTAL) a las mismas claves del plan de líneas."""
    return {
        "informes": kp["informes"],
        "finalizados": kp["finalizados"],
        "pend_elaborar": kp["pend_elaborar"],
        "lineas": kp["puntos"],
        "lineas_retiradas": kp["retirados"],
        "pend_inspeccion": kp["pend_inspeccion"],
        "por_elaborar": kp["pend_elaboracion"],
        "en_proceso": kp["en_proceso"],
        "por_revisar": kp["por_revisar"],
        "revisados": kp["revisados"],
        "en_revision": kp["en_revision"],
        "valorizados": kp["valorizados"],
    }


def total(kl, ka):
    return {k: kl.get(k, 0) + ka.get(k, 0) for k in set(kl) | set(ka)}


def avance(k):
    return (k["finalizados"] / k["informes"] * 100) if k["informes"] else 0.0


def comparativo(kl, ka):
    kt = total(kl, ka)
    filas = [{"INDICADOR": et, "PLAN DE LÍNEAS": kl[c], "PLANES ADICIONALES": ka[c],
              "TOTAL": kt[c]} for c, et in FILAS]
    filas.append({"INDICADOR": "% avance (finalizados)",
                  "PLAN DE LÍNEAS": f"{avance(kl):.1f}%",
                  "PLANES ADICIONALES": f"{avance(ka):.1f}%",
                  "TOTAL": f"{avance(kt):.1f}%"})
    tabla = pd.DataFrame(filas)
    tabla.index = range(1, len(tabla) + 1)
    return tabla


def alertas(kl, ka):
    """Pendientes críticos (solo los mayores que cero), de mayor a menor."""
    filas = []
    for c, et in ALERTAS:
        for plan, k in (("Plan de líneas", kl), ("Planes adicionales", ka)):
            if k[c] > 0:
                filas.append({"PLAN": plan, "PENDIENTE": et, "INFORMES": k[c]})
    filas.sort(key=lambda f: -f["INFORMES"])
    tabla = pd.DataFrame(filas, columns=["PLAN", "PENDIENTE", "INFORMES"])
    tabla.index = range(1, len(tabla) + 1)
    return tabla


def excel_resumen(kl, ka):
    hojas = {"RESUMEN": comparativo(kl, ka), "ALERTAS": alertas(kl, ka)}
    salida = io.BytesIO()
    with pd.ExcelWriter(salida, engine="openpyxl") as w:
        for nombre, tabla in hojas.items():
            tabla.to_excel(w, index=False, sheet_name=nombre)
            hoja = w.book[nombre]
            hoja.freeze_panes = "A2"
            relleno = PatternFill("solid", fgColor="0E2A47")
            for c in hoja[1]:
                c.font = Font(color="FFFFFF", bold=True)
                c.fill = relleno
                c.alignment = Alignment(horizontal="center", vertical="center",
                                        wrap_text=True)
            for col in hoja.columns:
                ancho = max(len(str(c.value if c.value is not None else "")) for c in col)
                hoja.column_dimensions[get_column_letter(col[0].column)].width = \
                    min(max(ancho + 2, 12), 40)
    salida.seek(0)
    return salida.getvalue()
